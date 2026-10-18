"""Per-source parsers. Each parse(path, fname) yields dicts with keys:
institution, person_name, title, department, base_pay, other_pay, total_pay, pay_components, fiscal_year, raw.
"""
import csv
import datetime as dt
import io
import json
import re
import zipfile

import openpyxl

from common import join_name, to_num

csv.field_size_limit(10_000_000)


def _k(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def pick(row, *cands):
    """Value for the first candidate header present (header match ignores case/punctuation)."""
    nk = {_k(k): v for k, v in row.items()}
    for c in cands:
        v = nk.get(_k(c))
        if v not in (None, ""):
            return v
    return None


def nsum(*xs):
    vals = [x for x in xs if x is not None]
    return sum(vals) if vals else None


def jsonl(path):
    with open(path) as f:
        for line in f:
            yield json.loads(line)


def xlsx_rows(path, header_match, sheet=None, max_scan=60):
    """Yield dict rows from the first sheet (or named one) whose header row contains header_match (case-insens.)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = [wb[sheet]] if sheet else wb.worksheets
    for ws in sheets:
        hdr = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if hdr is None:
                cells = [str(c).strip() if c is not None else "" for c in row]
                if any(_k(header_match) == _k(c) for c in cells):
                    hdr = cells
                elif i > max_scan:
                    break
                continue
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            d = {}
            for h, v in zip(hdr, row):
                if h:
                    if isinstance(v, str):
                        v = v.strip()
                    d[h] = v
            yield ws.title, d
        if hdr is not None and not sheet:
            return


_SMALL = {"of", "and", "the", "at", "for", "in", "on", "de"}


def tcase(s):
    """Title-case an all-caps institution name, keeping small words lower and & intact."""
    if not s:
        return s
    words = str(s).strip().lower().split()
    out = []
    for i, w in enumerate(words):
        if i and w in _SMALL:
            out.append(w)
        else:
            out.append(w[:1].upper() + w[1:])
    return " ".join(out).replace("Dupage", "DuPage").replace("Mchenry", "McHenry")


def comps(d):
    return {k: v for k, v in d.items() if v is not None}


# ---------------------------------------------------------------- Connecticut
CT_NAMES = {
    "University of Connecticut": "University of Connecticut",
    "UConn Health Center": "UConn Health",
    "CSU - Central": "Central Connecticut State University",
    "CSU - Eastern": "Eastern Connecticut State University",
    "CSU - Southern": "Southern Connecticut State University",
    "CSU - Western": "Western Connecticut State University",
    "Connecticut State Universities": "Connecticut State University System",
    "CT Community Colleges": "Connecticut Community Colleges",
    "Comm. College System Office": "Connecticut Community College System Office",
    "Charter Oak State College": "Charter Oak State College",
    "Connecticut Board of Regents": "Connecticut State Colleges & Universities (Board of Regents)",
}


CT_BOR_DEPT = {"Central CSU": "Central Connecticut State University", "Eastern CSU": "Eastern Connecticut State University",
               "Southern CSU": "Southern Connecticut State University", "Western CSU": "Western Connecticut State University",
               "Bd State Acdmc Awds": "Charter Oak State College", "CSCU System Office": "CSCU System Office",
               "System Office CSU": "CSCU System Office", "Board of Regents": "CSCU System Office",
               "CSCC Admin Office": "Connecticut State Community College"}


def parse_ct(path, fname):
    for r in jsonl(path):
        ag = r.get("agency")
        inst = CT_NAMES.get(ag) or (ag.replace("CCC - ", "") + " Community College" if ag.startswith("CCC - ") else ag)
        dep = r.get("department") or ""
        if ag == "Connecticut Board of Regents" and dep:
            inst = CT_BOR_DEPT.get(dep) or ("Connecticut State Community College" if dep.startswith("CSCC ") else inst)
        base, ot, oth, tot = (to_num(r.get(k)) for k in ("salaries_wages", "overtime", "other", "tot_gross"))
        yield dict(
            institution=inst,
            person_name=join_name(r.get("first_name"), r.get("middle_initial"), r.get("last_name"), r.get("name_suffix")),
            title=r.get("job_cd_descr"), department=r.get("department"),
            base_pay=base, other_pay=nsum(ot, oth), total_pay=tot,
            pay_components=comps({"salaries_wages": base, "overtime": ot, "other": oth, "tot_gross": tot,
                                  "annual_rate": to_num(r.get("annual_rate")), "paychecks": to_num(r.get("checks"))}),
            fiscal_year=int(r["pyrl_fiscal_yr"]), raw=r)


CT_CC = {"783": "Manchester", "785": "Northwestern Connecticut", "787": "Norwalk", "789": "Housatonic", "791": "Middlesex",
         "793": "Capital", "795": "Naugatuck Valley", "797": "Gateway", "799": "Tunxis", "801": "Three Rivers",
         "803": "Quinebaug Valley", "805": "Asnuntuck"}
CT_CC_AGENCY = {"CCC - " + v.replace(" Connecticut", ""): k for k, v in CT_CC.items()}
CT_CODE = {"777": "CSCU System Office", "778": "Charter Oak State College", "781": "Connecticut Community College System Office",
           "835": "Connecticut State University System Office", "840": "Central Connecticut State University",
           "845": "Western Connecticut State University", "850": "Southern Connecticut State University",
           "855": "Eastern Connecticut State University"}
CT_AGENCY2 = {"University of Connecticut": "University of Connecticut", "UConn Health Center": "UConn Health",
              "Charter Oak State College": "Charter Oak State College", "Board for Academic Awards": "Charter Oak State College",
              "Comm. College System Office": "Connecticut Community College System Office",
              "CSU - Central": "Central Connecticut State University", "CSU - Eastern": "Eastern Connecticut State University",
              "CSU - Southern": "Southern Connecticut State University", "CSU - Western": "Western Connecticut State University"}
CT_STATE_CC = "Connecticut State Community College"


def _ct_institution(r):
    """Institution + campus for one aggregated CT row (ct2 files).
    Before July 2023 each community college was its own institution; the 12 merged into Connecticut State Community
    College (CT State), whose payroll moved to the Board of Regents agency during FY2024 with campus in `department`.
    FY2010-2014 payroll used system agencies (CT Community Colleges, Connecticut State Universities), split here by
    the campus digits of deptid (e.g. CCC79700 -> 797 = Gateway), which match the per-campus agencies of FY2015-2023."""
    ag = r.get("agency") or ""
    dep = r.get("department") or ""
    code = r.get("campus_cd") or (r.get("deptid_max") or "")[3:6]
    if ag in CT_CC_AGENCY:
        return CT_CC[CT_CC_AGENCY[ag]] + " Community College", CT_CC[CT_CC_AGENCY[ag]]
    if ag in CT_AGENCY2:
        return CT_AGENCY2[ag], None
    if ag == "CT Community Colleges":
        if code in CT_CC:
            return CT_CC[code] + " Community College", CT_CC[code]
        return CT_CODE.get(code, "Connecticut Community Colleges"), None
    if ag == "Connecticut State Universities":
        return CT_CODE.get(code, "Connecticut State University System"), None
    if ag == "Connecticut Board of Regents":
        if dep.startswith("CSCC "):
            camp = dep[5:].replace(" Campus", "").strip()
            return CT_STATE_CC, (None if camp == "Admin Office" else camp)
        if dep in CT_BOR_DEPT:
            return CT_BOR_DEPT[dep], None
        if code in CT_CC:  # campus rows without a department value (post-merger agency)
            return CT_STATE_CC, CT_CC[code]
        return CT_CODE.get(code, "CSCU System Office"), None
    return ag, None


def parse_ct2(path, fname):
    rows = list(jsonl(path))
    fy = int(re.search(r"fy(\d{4})", fname).group(1))
    last = max((r.get("last_check") or "" for r in rows), default="")
    partial = last[:10] < f"{fy}-06-15"  # the payroll fiscal year ends June 30; files fetched mid-year are partial
    for r in rows:
        inst, campus = _ct_institution(r)
        base, ot, oth, tot = (to_num(r.get(k)) for k in ("salaries_wages", "overtime", "other", "tot_gross"))
        yield dict(
            institution=inst,
            person_name=join_name(r.get("first_name"), r.get("middle_initial"), r.get("last_name"), r.get("name_suffix")),
            title=r.get("job_cd_descr"), department=r.get("department"),
            base_pay=base, other_pay=nsum(ot, oth), total_pay=tot,
            pay_components=comps({"salaries_wages": base, "overtime": ot, "other": oth, "tot_gross": tot,
                                  "annual_rate": to_num(r.get("annual_rate")), "paychecks": to_num(r.get("checks")),
                                  "campus": campus, "campus_code": r.get("campus_cd") or (r.get("deptid_max") or "")[3:6] or None,
                                  "city": r.get("city_max"),
                                  "partial_year_through": last[:10] if partial else None}),
            fiscal_year=fy, raw=r)


# ---------------------------------------------------------------- Massachusetts
UMASS_ZIP = {"01003": "University of Massachusetts Amherst", "02125": "University of Massachusetts Boston",
             "02747": "University of Massachusetts Dartmouth", "01854": "University of Massachusetts Lowell",
             "01655": "UMass Chan Medical School", "01605": "UMass Chan Medical School",
             # the 'zip' field actually holds campus text in CTHRU
             "u mass amherst": "University of Massachusetts Amherst", "u mass boston": "University of Massachusetts Boston",
             "u mass dartmouth": "University of Massachusetts Dartmouth", "u mass lowell": "University of Massachusetts Lowell",
             "u mass worcester": "UMass Chan Medical School", "u mass central office": "University of Massachusetts System Office"}


def parse_ma(path, fname):
    for r in jsonl(path):
        div = re.sub(r"\s*\([A-Z]{2,4}\)\s*$", "", r.get("department_division") or "").strip()
        inst = tcase(div)
        if r.get("department_code") == "UMS":
            z = (r.get("department_location_zip_code") or "").strip()
            inst = UMASS_ZIP.get(z.lower()) or UMASS_ZIP.get(z[:5]) or "University of Massachusetts System"
        tot = to_num(r.get("pay_total_actual"))
        base = to_num(r.get("pay_base_actual"))
        pc = {k: to_num(v) for k, v in r.items() if (k.startswith("pay_") or re.fullmatch(r"a[a0-9]\w", k) or k == "annual_rate")}
        yield dict(
            institution=inst, person_name=join_name(r.get("name_first"), None, r.get("name_last")),
            title=r.get("position_title"), department=r.get("department_division"),
            base_pay=base, other_pay=(tot - base) if tot is not None and base is not None else None, total_pay=tot,
            pay_components=comps(pc), fiscal_year=int(r["year"]), raw=r)


# ---------------------------------------------------------------- NYC (CUNY community colleges)
def parse_nyc(path, fname):
    for r in jsonl(path):
        ag = r.get("agency_name", "")
        m = re.match(r"COMMUNITY COLLEGE \((.+)\)", ag)
        inst = (tcase(m.group(1)) + " Community College") if m else tcase(ag)
        inst = inst.replace("Laguardia", "LaGuardia").replace("Manhattan Community College", "Borough of Manhattan Community College")
        reg, ot, oth = (to_num(r.get(k)) for k in ("regular_gross_paid", "total_ot_paid", "total_other_pay"))
        yield dict(
            institution=inst, person_name=join_name(r.get("first_name"), r.get("mid_init"), r.get("last_name")),
            title=r.get("title_description"), department=r.get("work_location_borough"),
            base_pay=reg, other_pay=nsum(ot, oth), total_pay=nsum(reg, ot, oth),
            pay_components=comps({"base_salary_rate": to_num(r.get("base_salary")), "pay_basis": r.get("pay_basis"),
                                  "regular_gross_paid": reg, "total_ot_paid": ot, "total_other_pay": oth,
                                  "regular_hours": to_num(r.get("regular_hours")), "ot_hours": to_num(r.get("ot_hours"))}),
            fiscal_year=int(r["fiscal_year"]), raw=r)


# ---------------------------------------------------------------- Georgia
GA_RE = re.compile(r"UNIVERSITY|COLLEGE|INSTITUTE OF TECHNOLOGY|REGENTS OF THE UNIVERSITY", re.I)


def parse_ga(path, fname):
    z = zipfile.ZipFile(path)
    with z.open(z.namelist()[0]) as f:
        rd = csv.reader(io.TextIOWrapper(f, encoding="latin-1"))
        hdr = next(rd)
        for row in rd:
            if len(row) != len(hdr):
                continue
            r = dict(zip(hdr, [c.strip().strip("'") for c in row]))
            if not GA_RE.search(r["ORGANIZATION"]):
                continue
            sal = to_num(r["SALARY"])
            yield dict(institution=tcase(r["ORGANIZATION"]),
                       person_name=r["NAME"], title=r["TITLE"], department=None,
                       base_pay=None, other_pay=None, total_pay=sal,
                       pay_components=comps({"salary": sal, "travel": to_num(r["TRAVEL"])}),
                       fiscal_year=int(r["FISCAL_YEAR"]), raw=r)


# ---------------------------------------------------------------- Texas College Salaries
def parse_tx(path, fname):
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        for r in csv.DictReader(f):
            fy = None
            try:
                d = dt.datetime.strptime(r["data_date"], "%m/%d/%Y")
                fy = d.year + (1 if d.month >= 9 else 0)  # Texas state FY runs Sep-Aug
            except (ValueError, KeyError):
                pass
            sal = to_num(r.get("salary"))
            yield dict(institution=r["institution"], person_name=r["full_name"], title=r.get("job_title"),
                       department=r.get("department"), base_pay=sal, other_pay=None, total_pay=None,
                       pay_components=comps({"salary": sal, "employment_time": r.get("employment_time"), "data_date": r.get("data_date")}),
                       fiscal_year=fy, raw=r)


# ---------------------------------------------------------------- Illinois community colleges (ICCB)
ICCB_FIX = {"Ccc District Office": "City Colleges of Chicago District Office", "Kennedy-king College": "Kennedy-King College",
            "Olive-harvey College": "Olive-Harvey College", "Il Eastern District Office": "Illinois Eastern Community Colleges District Office"}


def parse_il_iccb(path, fname):
    for _, r in xlsx_rows(path, "Last Name"):
        base = to_num(pick(r, "Base Pay"))
        extra = to_num(pick(r, "Extra Duty Pay"))
        enh = to_num(pick(r, "Enhancements"))
        pc = {k: (to_num(v) if to_num(v) is not None else v) for k, v in r.items()
              if _k(k) in ("basepay", "benefits", "annuities", "enhancements", "extradutypay", "vacationdays", "sickdays", "personaldays")}
        fy = pick(r, "Fiscal Year")
        inst = tcase(pick(r, "College Name") or "")
        yield dict(institution=ICCB_FIX.get(inst, inst),
                   person_name=join_name(pick(r, "First Name"), pick(r, "Middle Initial"), pick(r, "Last Name")),
                   title=pick(r, "Title"), department=pick(r, "Employment Classification"),
                   base_pay=base, other_pay=nsum(extra, enh), total_pay=nsum(base, extra, enh),
                   pay_components=comps(pc), fiscal_year=int(str(fy)[:4]) if fy else int(re.search(r"fy(\d{4})", fname).group(1)),
                   raw=r)


# ---------------------------------------------------------------- Ohio State (calendar-year earnings files)
def parse_oh_osu(path, fname):
    year = int(re.search(r"osu_(\d{4})_", fname).group(1))
    for _, r in xlsx_rows(path, "Gross Pay"):
        last = pick(r, "Last Name - Legal", "Last Name", "Legal Last Name", "Preferred Last Name")
        first = pick(r, "First Name - Preferred", "First Name", "Preferred First Name")
        reg = to_num(pick(r, "Regular Pay (Base Pay + Paid Time Off + Premium)", "Regular Pay", "Regular"))
        bonus, ot = to_num(pick(r, "Bonus")), to_num(pick(r, "Overtime"))
        oth = to_num(pick(r, "Other (Allowance + Supplemental + Uncategorized)", "Other"))
        gross = to_num(pick(r, "Gross Pay"))
        yield dict(institution="The Ohio State University", person_name=join_name(first, None, last),
                   title=pick(r, "Job Profile Name", "Job Profile", "Job Title"),
                   department=pick(r, "Cost Center", "Department"),
                   base_pay=reg, other_pay=nsum(bonus, ot, oth), total_pay=gross,
                   pay_components=comps({"regular": reg, "bonus": bonus, "overtime": ot, "other": oth, "gross": gross,
                                         "position_group": pick(r, "Position Group")}),
                   fiscal_year=year, raw=r)


def parse_oh_ou(path, fname):
    for _, r in xlsx_rows(path, "LAST_NAME"):
        base = to_num(pick(r, "BASE_SALARY"))
        addl = to_num(pick(r, "ADDL_COMP"))
        w2 = to_num(pick(r, "MEDICARE WAGES"))
        yield dict(institution="Ohio University", person_name=join_name(pick(r, "FIRST_NAME"), pick(r, "MIDDLE_NAME"), pick(r, "LAST_NAME")),
                   title=pick(r, "TITLE", "ACADEMIC_TITLE"), department=pick(r, "DEPARTMENT_NAME"),
                   base_pay=base, other_pay=addl, total_pay=w2,
                   pay_components=comps({"base_salary": base, "addl_comp": addl, "medicare_wages_2025": w2,
                                         "pay_basis": pick(r, "PAY_BASIS"), "campus": pick(r, "LOCATION_NAME")}),
                   fiscal_year=2025, raw=r)


# ---------------------------------------------------------------- Minnesota State (MMB salary files)
# MMB location names -> Minnesota State college/university (substring match, first hit wins). Colleges keep the names
# they had at the time; Minnesota North College (2022 merger of Hibbing, Itasca, Mesabi Range, Rainy River and
# Vermilion) keeps the campus college name so the series stays continuous.
MN_CAMPUS = [
    ("St Cloud SU", "St. Cloud State University"), ("SCSU", "St. Cloud State University"),
    ("Mankato SU", "Minnesota State University, Mankato"), ("Metro SU", "Metropolitan State University"),
    ("Winona SU", "Winona State University"), ("Moorhead SU", "Minnesota State University Moorhead"),
    ("Bemidji SU", "Bemidji State University"), ("Southwest SU", "Southwest Minnesota State University"),
    ("Century", "Century College"), ("Normandale", "Normandale Community College"),
    ("S Cntrl", "South Central College"), ("St Cloud TC", "St. Cloud Technical & Community College"),
    ("Minneapolis CC", "Minneapolis College"), ("Minneapolis TC", "Minneapolis College"),
    ("Inver Hills", "Inver Hills Community College"), ("Anoka Ramsey", "Anoka-Ramsey Community College"),
    ("Anoka/Ram", "Anoka-Ramsey Community College"), ("Anoka Hennepin", "Anoka Technical College"),
    ("St Paul TC", "Saint Paul College"), ("North Hennepin", "North Hennepin Community College"),
    ("Hennepin North", "Hennepin Technical College"), ("Hennepin South", "Hennepin Technical College"),
    ("Hennepin Dist", "Hennepin Technical College"), ("Lake Superior", "Lake Superior College"),
    ("Alexandria", "Alexandria Technical & Community College"), ("Rochester", "Rochester Community and Technical College"),
    ("Dakota", "Dakota County Technical College"), ("Mn State CTC", "Minnesota State Community and Technical College"),
    ("Cntrl Lakes", "Central Lakes College"), ("Northland", "Northland Community and Technical College"),
    ("Riverland", "Riverland Community College"), ("Fond du Lac", "Fond du Lac Tribal and Community College"),
    ("Ridgewater", "Ridgewater College"), ("Itasca", "Itasca Community College"), ("Hibbing", "Hibbing Community College"),
    ("Mesabi", "Mesabi Range College"), ("Vermilion", "Vermilion Community College"), ("Rainy River", "Rainy River Community College"),
    ("Arrowhead", "Northeast Higher Education District"), ("Southeast", "Minnesota State College Southeast"),
    ("Pine TC", "Pine Technical & Community College"), ("MnWest", "Minnesota West Community and Technical College"),
    ("Northwest Bemidji", "Northwest Technical College"), ("System Office", "Minnesota State System Office"),
]
MN_SYSTEM = "Minnesota State Colleges and Universities"


def _mn_campus(loc):
    loc = str(loc or "").lower()
    for pat, inst in MN_CAMPUS:
        if pat.lower() in loc:
            return inst
    return None


def parse_mn(path, fname):
    fy = int(re.search(r"fy(\d{4})", fname).group(1))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    hr_ws = [w for w in wb.worksheets if "HR" in w.title.upper()][0]
    er_ws = [w for w in wb.worksheets if "EARN" in w.title.upper()][0]
    it = er_ws.iter_rows(values_only=True)
    eh = [str(h) for h in next(it)]
    earn = {}
    for row in it:
        d = dict(zip(eh, row))
        earn[str(d[eh[0]])] = d
    it = hr_ws.iter_rows(values_only=True)
    hh = [str(h) for h in next(it)]
    people = {}
    dept_campus = {}  # department number -> {campus: rows} from rows whose location names a campus
    for row in it:
        d = dict(zip(hh, row))
        ag = pick(d, "AGENCY_NAME", "EMPLT_AGENCY_NM") or ""
        if "Colleges & Universities" not in ag and "MnSCU" not in ag:
            continue
        camp = _mn_campus(pick(d, "LOCATION_NAME", "LOCN_NM"))
        if camp:
            dc = dept_campus.setdefault(str(pick(d, "DEPARTMENT_NBR", "DEPT_NBR")), {})
            dc[camp] = dc.get(camp, 0) + 1
        tid = str(pick(d, "TEMPORARY_ID", "TEMPORARY_EMPLOYEE_ID"))
        rec = to_num(pick(d, "RECORD_NBR", "EMPL_REC_NBR")) or 0
        # one row per person: keep the lowest employment record (primary appointment); others go in pay_components
        if tid in people:
            people[tid]["other_positions"].append(pick(d, "JOB_TITLE", "JOB_DESC"))
            if rec >= people[tid]["rec"]:
                continue
            people[tid].update(rec=rec, hr=d)
        else:
            people[tid] = {"rec": rec, "hr": d, "other_positions": []}
    for tid, p in people.items():
        d = p["hr"]
        e = earn.get(tid, {})
        loc = pick(d, "LOCATION_NAME", "LOCN_NM")
        camp, how = _mn_campus(loc), "location"
        if not camp:
            dc = dept_campus.get(str(pick(d, "DEPARTMENT_NBR", "DEPT_NBR")))
            if dc:
                camp, how = max(dc.items(), key=lambda kv: kv[1])[0], "department number (majority location)"
        reg = to_num(pick(e, "REGULAR_WAGES", "REGWAGES", "SALARY"))
        ot = to_num(pick(e, "OVERTIME_WAGES", "OTWAGES", "OVERTIME"))
        oth = to_num(pick(e, "OTHER_WAGES")) if pick(e, "OTHER_WAGES") is not None else nsum(to_num(pick(e, "DIFFERENTIALS_ETC")), to_num(pick(e, "OTHER_PAY")))
        tot = to_num(pick(e, "TOTAL_WAGES", "TOTALWAGES", "TOTAL_PAY")) if pick(e, "TOTAL_WAGES", "TOTALWAGES", "TOTAL_PAY") is not None else nsum(reg, ot, oth)
        dept = pick(d, "DEPARTMENT_NAME", "DEPT_NM")
        yield dict(institution=camp or MN_SYSTEM, person_name=pick(d, "EMPLOYEE_NAME", "EMPL_NM"),
                   title=pick(d, "JOB_TITLE", "JOB_DESC"), department=dept,
                   base_pay=reg, other_pay=nsum(ot, oth), total_pay=tot,
                   pay_components=comps({**{k: v for k, v in e.items() if k != eh[0]},
                                         "compensation_rate": to_num(pick(d, "COMPENSATION_RATE", "CMPTN_RT")),
                                         "comp_frequency": pick(d, "COMP_FREQUENCY_DESC", "CMPTN_FREQ_DESC"),
                                         "position_fte": to_num(pick(d, "POSITION_FTE", "POSN_FTE_PRCT_RT")),
                                         "location": loc, "campus_from": how if camp else None,
                                         "other_positions": p["other_positions"] or None,
                                         "note": "earnings are FY totals across all state positions"}),
                   fiscal_year=fy, raw={"hr": d, "earnings": e})


# ---------------------------------------------------------------- Washington (data.wa.gov 2010-2013, wide)
def _wa_higher(code):
    c = int(code)
    return c == 352 or 360 <= c <= 380 or 600 <= c <= 699


def parse_wa(path, fname):
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        for r in csv.DictReader(f):
            if not _wa_higher(r["Agency"]):
                continue
            for y in (2010, 2011, 2012, 2013):
                v = to_num(r.get(f"Salary{y}"))
                if not v:
                    continue
                yield dict(institution=r["AgencyTitle"], person_name=r["EmployeeName"], title=r["JobTitle"], department=None,
                           base_pay=None, other_pay=None, total_pay=v, pay_components={"salary": v, "calendar_year": y},
                           fiscal_year=y, raw=r)


# ---------------------------------------------------------------- Kansas (Kansas Policy Institute compiled payroll)
KS_RE = re.compile(r"University|Medical Center", re.I)


def parse_ks(path, fname):
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        for r in csv.DictReader(f):
            ag = r["Agency"]
            if not KS_RE.search(ag) or "Emergency" in ag:
                continue
            inst = {"KUMC University of Kansas Medical Center": "University of Kansas Medical Center",
                    "Univ of Ks Medical Center": "University of Kansas Medical Center",
                    "Pittsburgh State University": "Pittsburg State University"}.get(ag, ag)
            tot = to_num(r["TotalPay"])
            yield dict(institution=inst, person_name=join_name(r["FirstName"], None, r["LastName"]), title=r["Position"],
                       department=None, base_pay=None, other_pay=None, total_pay=tot, pay_components={"total_pay": tot},
                       fiscal_year=int(r["Year"]), raw=r)


# ---------------------------------------------------------------- Nebraska (budgeted salaries)
NE_CAMPUS = {"UNL": "University of Nebraska-Lincoln", "UNO": "University of Nebraska Omaha", "UNMC": "University of Nebraska Medical Center",
             "UNK": "University of Nebraska at Kearney", "NCTA": "Nebraska College of Technical Agriculture",
             "CS": "University of Nebraska System (Central Administration)", "UNCA": "University of Nebraska System (Central Administration)"}


def parse_ne(path, fname):
    y = re.search(r"ne_(\d{4})_(\d{4})", fname)
    fy = int(y.group(2))
    for _, r in xlsx_rows(path, "Employee"):
        name = pick(r, "Employee")
        if not name or str(name).lower().startswith("total"):
            continue
        sal = to_num(pick(r, "Budgeted Annual Salary", "Budgeted \nAnnual\nSalary"))
        if sal is None:
            sal = to_num(next((v for k, v in r.items() if "budgeted" in k.lower()), None))
        st = to_num(next((v for k, v in r.items() if "state" in k.lower()), None))
        ot = to_num(next((v for k, v in r.items() if "other" in k.lower()), None))
        campus = pick(r, "Campus")
        yield dict(institution=NE_CAMPUS.get(str(campus).strip(), f"University of Nebraska ({campus})"), person_name=name,
                   title=pick(r, "Position"), department=pick(r, "Department"), base_pay=sal, other_pay=None, total_pay=None,
                   pay_components=comps({"budgeted_annual_salary": sal, "salary_state_aided_funds": st, "salary_other_funds": ot, "campus": campus}),
                   fiscal_year=fy, raw={k.replace("\n", " "): v for k, v in r.items()})


# ---------------------------------------------------------------- Maine
ME_CAMPUS = {"MLAW": "University of Maine School of Law", "SWS": "University of Maine System (System Wide Services)",
             "UM": "University of Maine", "UMA": "University of Maine at Augusta", "UMF": "University of Maine at Farmington",
             "UMFK": "University of Maine at Fort Kent", "UMM": "University of Maine at Machias",
             "UMPI": "University of Maine at Presque Isle", "USM": "University of Southern Maine"}


def parse_me(path, fname):
    for _, r in xlsx_rows(path, "Formatted Name"):
        sal = to_num(pick(r, "Salbase"))
        c = pick(r, "Campus")
        yield dict(institution=ME_CAMPUS.get(c, c), person_name=pick(r, "Formatted Name"), title=pick(r, "Title"),
                   department=pick(r, "Dept ID"), base_pay=sal, other_pay=None, total_pay=None,
                   pay_components=comps({"salbase": sal, "fte": to_num(pick(r, "FTE")), "job_status": pick(r, "JobSt"),
                                         "faculty_appt": pick(r, "Faculty Appt")}),
                   fiscal_year=2026, raw=r)


# ---------------------------------------------------------------- New Mexico (UNM sunshine salary book)
def parse_nm(path, fname):
    data = json.load(open(path))["data"]
    cols = ["name", "department", "col3", "title", "position", "fte_pct", "salary"]
    for row in data:
        r = dict(zip(cols, row))
        sal = to_num(r.get("salary"))
        yield dict(institution="University of New Mexico", person_name=r["name"], title=r.get("title"), department=r.get("department"),
                   base_pay=sal, other_pay=None, total_pay=None,
                   pay_components=comps({"salary": sal, "fte_pct": to_num(r.get("fte_pct")), "position": r.get("position")}),
                   fiscal_year=2027, raw={"row": row})


# ---------------------------------------------------------------- DC (UDC rows of DC employee layer)
def parse_dc(path, fname):
    for r in jsonl(path):
        rate = to_num(r.get("COMPRATE"))
        annual = rate if rate is not None and rate >= 1000 else None
        yield dict(institution="University of the District of Columbia", person_name=join_name(r.get("FIRST_NAME"), None, r.get("LAST_NAME")),
                   title=r.get("JOBTITLE"), department=None, base_pay=annual, other_pay=None, total_pay=None,
                   pay_components=comps({"comprate": rate, "rate_kind": "annual" if annual else "hourly/per-unit", "grade": r.get("GRADE"),
                                         "appointment": r.get("GVT_TYPE_OF_APPT")}),
                   fiscal_year=2027, raw=r)


# ---------------------------------------------------------------- Arizona
def parse_az_asu(path, fname):
    fy_file = int(re.search(r"(\d{4})", fname).group(1))
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
        for r in csv.DictReader(f):
            name = pick(r, "Full Name", "Name")
            if not name:
                continue
            sal = to_num(pick(r, "Salary", "Total Salary"))
            fte = to_num(pick(r, "FTE", "Fte", "Total FTE", "Total Fte"))
            yield dict(institution="Arizona State University", person_name=name.rstrip(" -"),
                       title=pick(r, "Job Description", "Jobcode Descr", "Jobcode Description", "Title"),
                       department=pick(r, "Department Description", "Department Descr", "Department"),
                       base_pay=sal, other_pay=None, total_pay=None,
                       pay_components=comps({"salary": sal, "fte": fte, "college_unit": pick(r, "College/Unit")}),
                       fiscal_year=fy_file, raw=r)


def parse_az_ua(path, fname):
    html = open(path, encoding="utf-8", errors="replace").read()
    cols = json.loads(re.search(r"_Flourish_data_column_names\s*=\s*(\{.*?\}\}),", html).group(1))["rows"]["columns"]
    i = html.index("_Flourish_data = ") + len("_Flourish_data = ")
    data, _ = json.JSONDecoder().raw_decode(html[i:])
    for item in data["rows"]:
        r = dict(zip(cols, item["columns"]))
        sal = to_num(r.get("Annual Salary >> at Employment FTE"))
        yield dict(institution="University of Arizona", person_name=r.get("Name"), title=r.get("Primary Title"), department=r.get("Department"),
                   base_pay=sal, other_pay=None, total_pay=None,
                   pay_components=comps({"annual_salary_at_fte": sal, "full_fte_salary": to_num(r.get("Full FTE Salary")),
                                         "fte": to_num(r.get("FTE >> Full Time Equivalent")), "hourly_rate": to_num(r.get("Hourly Pay Rate >> Primary Job")),
                                         "ft_pt": r.get("FT/PT >> Full Time or Part Time")}),
                   fiscal_year=int(r["Fiscal Year"]), raw=r)


# ---------------------------------------------------------------- Oklahoma (OMES payroll, FY2026 monthly payment lines)
OK_RE = re.compile(r"UNIV|COLLEGE|COLEGE", re.I)


def parse_ok(path, fname):
    """One file per month; rows are payment lines. Aggregated per agency + last name + initials within the file (month)."""
    agg = {}
    with open(path, encoding="latin-1", newline="") as f:
        for r in csv.DictReader(f):
            if not OK_RE.search(r["AGENCY_NAME"]):
                continue
            k = (r["AGENCY_NUMBER"], r["AGENCY_NAME"], r["LAST_NAME"].strip(), r["FIRST_INITIAL"].strip(), r["MIDDLE_INITIAL"].strip())
            a = agg.setdefault(k, {"amount": 0.0, "hours": 0.0, "accounts": {}, "lines": 0, "report_date": r["REPORT_DATE"]})
            amt = to_num(r["AMOUNT"]) or 0.0
            a["amount"] += amt
            a["hours"] += to_num(r["HOURS"]) or 0.0
            a["accounts"][r["ACCOUNT_DESCRIPTION"]] = round(a["accounts"].get(r["ACCOUNT_DESCRIPTION"], 0.0) + amt, 2)
            a["lines"] += 1
    ym = re.search(r"(\d{4})(\d{2})", fname)
    fy = int(ym.group(1)) + (1 if int(ym.group(2)) >= 7 else 0)
    for (num, ag, last, fi, mi), a in agg.items():
        reg = sum(v for k2, v in a["accounts"].items() if "regular" in k2.lower())
        yield dict(institution=tcase(ag), person_name=join_name(fi, mi, last), title=None, department=None,
                   base_pay=round(reg, 2), other_pay=round(a["amount"] - reg, 2), total_pay=round(a["amount"], 2),
                   pay_components={"month": ym.group(0), "by_account": a["accounts"], "hours": round(a["hours"], 2), "payment_lines": a["lines"],
                                   "note": "monthly total; name is last name + initials only"},
                   fiscal_year=fy, raw={"agency_number": num, "agency_name": ag, "last_name": last, "first_initial": fi,
                                        "middle_initial": mi, "report_date": a["report_date"]})


# ---------------------------------------------------------------- New Jersey (pension membership, higher-ed employers)
NJ_EXCLUDE = re.compile(r"C/S|CHRTR|CHARTER|PREP|BORO|HOSPITAL|ACHIEVE", re.I)


def parse_nj(path, fname):
    for r in jsonl(path):
        emp = r.get("employer_name", "")
        if NJ_EXCLUDE.search(emp):
            continue
        sal = to_num(r.get("current_employer_salary_amt"))
        yield dict(institution=tcase(emp), person_name=join_name(r.get("member_first_name"), r.get("member_mi"), r.get("member_last_name")),
                   title=None, department=r.get("location_name"), base_pay=sal, other_pay=None, total_pay=None,
                   pay_components=comps({"current_employer_salary_amt": sal, "all_employers_salary_amt": to_num(r.get("all_employers_salary_amt")),
                                         "pension_fund": r.get("pension_fund_name")}),
                   fiscal_year=2026, raw=r)


# ---------------------------------------------------------------- University of Illinois Gray Book (HTML)
import html as _html

GB_CAMPUS = {"Urbana-Champaign": "University of Illinois Urbana-Champaign", "Chicago": "University of Illinois Chicago",
             "Springfield": "University of Illinois Springfield", "System": "University of Illinois System Offices"}
_TR = re.compile(r'<tr id="(c\d+-d\d+-e(\d+)-j(\d+))">(.*?)</tr>', re.S)
_TD = re.compile(r'<td headers="c\d+-d\d+-([a-z-]+)-header"[^>]*>(.*?)</td>', re.S)
_HDR = re.compile(r'<h([23]) id="(c\d+(?:-d\d+)?)">(.*?)</h[23]>', re.S)


def _txt(x):
    return _html.unescape(re.sub(r"<[^>]+>", " ", x)).replace("\xa0", " ").strip()


def parse_il_graybook(path, fname):
    s = open(path, encoding="utf-8", errors="replace").read()
    m = re.search(r'<span class="center">\(([^)]*)\)</span>', s)
    camp_txt = m.group(1) if m else ""
    inst = next((v for k, v in GB_CAMPUS.items() if k.lower() in camp_txt.lower()), "University of Illinois System")
    heads = {mm.group(2): _txt(mm.group(3)) for mm in _HDR.finditer(s)}
    last_name_by_emp = {}
    for mm in _TR.finditer(s):
        rid, emp, job, body = mm.group(1), mm.group(2), mm.group(3), mm.group(4)
        cells = {k: _txt(v) for k, v in _TD.findall(body)}
        c, d = rid.split("-")[0], "-".join(rid.split("-")[:2])
        ekey = (d, emp)
        name = cells.get("empl") or last_name_by_emp.get(ekey)
        if cells.get("empl"):
            last_name_by_emp[ekey] = cells["empl"]
        prop, pres = to_num(cells.get("propsal")), to_num(cells.get("pressal"))
        yield dict(institution=inst, person_name=name, title=cells.get("job"), department=heads.get(d),
                   base_pay=prop, other_pay=None, total_pay=None,
                   pay_components=comps({"proposed_salary_fy2027": prop, "present_salary_fy2026": pres,
                                         "proposed_fte": to_num(cells.get("propfte")), "present_fte": to_num(cells.get("presfte")),
                                         "tenure": cells.get("tenure") or None, "empl_class": cells.get("empclass"),
                                         "job_row": int(job), "college": heads.get(c)}),
                   fiscal_year=2027, raw={"row_id": rid, "page": fname, "campus": camp_txt, "college": heads.get(c),
                                          "department": heads.get(d), **cells})


# ---------------------------------------------------------------- South Carolina state salaries (HTML)
_SC_TR = re.compile(r"<tr>(.*?)</tr>", re.S)
_SC_TD = re.compile(r'<td headers="view-([a-z0-9-]+)-table-column"[^>]*>(.*?)</td>', re.S)


def parse_sc(path, fname):
    s = open(path, encoding="utf-8", errors="replace").read()
    for page in s.split("<!-- UNIPAY PAGE ")[1:]:
        body = page.split("<tbody", 1)[-1].split("</tbody>", 1)[0]
        for tr in _SC_TR.findall(body):
            c = {k: _txt(v) for k, v in _SC_TD.findall(tr)}
            if not c:
                continue
            tot = to_num(c.get("field-total-compensation"))
            bon = to_num(c.get("field-bonuses"))
            yield dict(institution=tcase(c.get("field-agency")).replace("Of Sc", "of SC").replace("of Sc", "of SC"),
                       person_name=join_name(c.get("field-first-name"), None, c.get("label")), title=c.get("field-job-title-sq"),
                       department=None, base_pay=None, other_pay=bon, total_pay=tot,
                       pay_components=comps({"total_compensation": tot, "bonuses": bon}),
                       fiscal_year=2027, raw=c)


# ---------------------------------------------------------------- Maryland (UMD Diamondback salary guide API)
def parse_md(path, fname):
    """One row per API row. Year label = fiscal year: checked against known tenure dates (Loh president through
    '2020', Pines from '2021' = July 2020 handover; Locksley absent 2016-2018 and head coach from '2019' = left Jan 2016,
    returned Dec 2018; Willard '2022'-'2024' = Mar 2022-Mar 2025). Departed coaches with settlements stay listed
    while paid (Durkin '2019'-'2021', Turgeon '2022'-'2025')."""
    y = int(re.search(r"md_(\d{4})", fname).group(1))
    for r in jsonl(path):
        sal = to_num(r.get("Salary"))
        name = " ".join(str(r.get("Employee") or "").split()) or None
        yield dict(institution="University of Maryland, College Park", person_name=name,
                   title=(r.get("Title") or "").strip() or None, department=(r.get("Department") or "").strip() or None,
                   base_pay=sal, other_pay=None, total_pay=None,
                   pay_components=comps({"salary": sal, "division": (r.get("Division") or "").replace("\n", " ").strip() or None,
                                         "category_status": (r.get("Category Status") or "").replace("\n", " ").strip() or None,
                                         "employee_type": r.get("Employee Type"), "exempt_status": r.get("Exempt Status"),
                                         "api_page": r.get("_page")}),
                   fiscal_year=y, raw=r)


# ---------------------------------------------------------------- Nebraska community colleges (Flatwater Free Press)
NE_CC_PAY = ("salary_pay", "overtime_pay", "bonus_pay", "additional_pay", "holiday_pay", "vacation_pay", "sick_pay", "comp_pay")


def parse_ne_cc(path, fname):
    for r in jsonl(path):
        a = r.get("acf") or {}
        org = _html.unescape(re.sub(r"<[^>]+>", "", (r.get("excerpt") or {}).get("rendered", ""))).split(" – ")[0].strip()
        pay = {k: to_num(a.get(k)) for k in NE_CC_PAY}
        base = pay["salary_pay"]
        other = nsum(*(v for k, v in pay.items() if k != "salary_pay"))
        tot = to_num(a.get("total_pay"))
        yr = str(a.get("year") or "").strip()
        yield dict(institution=org or None, person_name=join_name(a.get("first_name"), None, a.get("last_name")),
                   title=a.get("job_title") or None, department=a.get("department") or a.get("division") or None,
                   base_pay=base, other_pay=other, total_pay=tot,
                   pay_components=comps({**pay, "total_pay": tot, "pay_rate": to_num(a.get("pay_rate")), "wage_type": a.get("wage_type") or None,
                                         "ft_pt": a.get("ft_pt") or None, "hire_date": a.get("hire_date"), "end_date": a.get("end_date"),
                                         "ffp_update_batch": a.get("ffp_update_batch"), "flatwater_id": a.get("flatwater_id")}),
                   fiscal_year=int(yr[:4]) if yr[:4].isdigit() else None,
                   raw={"id": r.get("id"), "link": r.get("link"), "title": (r.get("title") or {}).get("rendered"),
                        "excerpt": (r.get("excerpt") or {}).get("rendered"), "organization": r.get("organization"), "acf": a})


# ---------------------------------------------------------------- St. Louis Post-Dispatch Public Pay (universities)
STL_INST = {"University of Missouri-Columbia": ("MO", "University of Missouri-Columbia"),
            "University of Missouri-Kansas City": ("MO", "University of Missouri-Kansas City"),
            "University of Missouri-St. Louis": ("MO", "University of Missouri-St. Louis"),
            "Missouri University of Science and Technology": ("MO", "Missouri University of Science and Technology"),
            "University of Missouri System": ("MO", "University of Missouri System Office"),
            "University of Missouri Health Care": ("MO", "University of Missouri Health Care"),
            "St. Louis Community College": ("MO", "St. Louis Community College"),
            "St. Charles Community College": ("MO", "St. Charles Community College"),
            "Southern Illinois University Edwardsville": ("IL", "Southern Illinois University Edwardsville"),
            "University of Illinois Administration": ("IL", "University of Illinois System Offices"),
            "University of Illinois-Chicago": ("IL", "University of Illinois Chicago"),
            "University of Illinois-Springfield": ("IL", "University of Illinois Springfield"),
            "University of Illinois-Urbana-Champaign": ("IL", "University of Illinois Urbana-Champaign")}
_STL_TR = re.compile(r'<tr>\s*<th class="str-col" scope="row">(.*?)</th>\s*<td class="str-col">(.*?)</td>\s*<td class="num-col">(.*?)</td>\s*<td class="num-col">(.*?)</td>', re.S)


def parse_mo_stl(path, fname, state=None):
    s = open(path, encoding="utf-8", errors="replace").read()
    agency = re.search(r"<!-- UNIPAY AGENCY \d+ (.*?) https?://", s).group(1).strip()
    st, inst = STL_INST[agency]
    if state and st != state:
        return
    for part in s.split("<!-- UNIPAY DEPT ")[1:]:
        m = re.match(r"(\S+) PAGE (\d+) NAME (.*?) URL (\S+) -->", part)
        did, page, dname, url = m.group(1), int(m.group(2)), _html.unescape(m.group(3)), m.group(4)
        body = part.split("<table", 1)[-1].split("</table>", 1)[0]
        for name, pos, pay, hired in _STL_TR.findall(body):
            name, pos, pay, hired = _txt(name), _txt(pos), _txt(pay), _txt(hired)
            tot = to_num(pay)
            yield dict(institution=inst, person_name=name or None, title=pos or None, department=dname or None,
                       base_pay=None, other_pay=None, total_pay=tot,
                       pay_components=comps({"gross_pay_2025": tot, "hire_date": None if hired in ("", "Not provided") else hired}),
                       fiscal_year=2025,
                       raw={"agency": agency, "dept_id": did, "dept": dname, "page": page, "url": url,
                            "name": name, "position": pos, "gross_pay": pay, "date_hired": hired})


def parse_mo_stl_mo(path, fname):
    return parse_mo_stl(path, fname, "MO")


def parse_mo_stl_il(path, fname):
    return parse_mo_stl(path, fname, "IL")


SPECS = {
    "ct": dict(state="CT", system="State of Connecticut payroll (higher-ed agencies)", catalog_url_like="%9m78-yc88%", parse=parse_ct2,
               file_filter=lambda f: f.startswith("ct2_"),
               notes="fiscal_year = CT payroll fiscal year (Jul-Jun); one row per employee record per agency (per campus code for system agencies) per FY, "
                     "aggregated server-side; campus from deptid digits; community colleges are 'Connecticut State Community College' after the "
                     "July 2023 merger (campus in pay_components.campus); the latest FY is partial (pay_components.partial_year_through)"),
    "ma": dict(state="MA", system="Massachusetts CTHRU statewide payroll (higher-ed departments)", catalog_url_like="%9ttk-7vz6%", parse=parse_ma,
               notes="fiscal_year = calendar year (CTHRU 'year'); UMass campus inferred from department_location_zip_code"),
    "nyc": dict(state="NY", system="NYC Citywide Payroll (CUNY community colleges)", catalog_url_like="%k397-673e%", parse=parse_nyc,
                notes="fiscal_year = NYC FY (Jul-Jun)"),
    "ga": dict(state="GA", system="Open Georgia salaries (USG, TCSG, GMC)", catalog_url_like="%open.ga.gov%", parse=parse_ga,
               notes="filtered ORGANIZATION to universities/colleges; salary only (travel kept in pay_components)"),
    "tx": dict(state="TX", system="Texas College Salaries (volunteer project, TPIA requests)", catalog_url_like="%drive.google.com%", parse=parse_tx,
               notes="bulk export publishes first name + last initial only; base salary snapshot; fiscal_year derived from data_date (Sep-Aug FY)"),
    "il_iccb": dict(state="IL", system="Illinois Community College Board salary & benefits", catalog_url_like="%iccb%", parse=parse_il_iccb,
                    notes="total_pay = base + extra duty + enhancements; benefits/annuities kept in pay_components only"),
    "oh_osu": dict(state="OH", system="The Ohio State University", catalog_url_like="%osu.edu%", parse=parse_oh_osu,
                   file_filter=lambda f: "earnings" in f,
                   notes="calendar-year paid earnings files; salary-snapshot files downloaded but not loaded (same people, rate not pay)"),
    "oh_ou": dict(state="OH", system="Ohio University", catalog_url_like="%ohio.edu%", parse=parse_oh_ou,
                  notes="base salary as of 12/31/2025 plus 2025 Medicare (W-2) wages as total_pay"),
    "mn": dict(state="MN", system="Minnesota MMB state employee salary data (Minnesota State only)", catalog_url_like="%mn.gov/mmb%", parse=parse_mn,
               notes="agency 'MN St Colleges & Universities'; one row per person per FY (primary record); earnings are all state positions; "
                     "institution = campus from the MMB location name (department-number majority when the location is blank)"),
    "wa": dict(state="WA", system="Washington state salaries 2010-2013 (data.wa.gov y3ds-rkew)", catalog_url_like="%y3ds-rkew%", parse=parse_wa,
               notes="wide file unpivoted to one row per person/agency/calendar year with non-zero salary; higher-ed agency codes 352, 360-380, 600-699"),
    "ks": dict(state="KS", system="Kansas state payroll 2021-2025 (Kansas Policy Institute / KansasOpenGov)", catalog_url_like="%kansasopengov%", parse=parse_ks,
               notes="third-party compiled from KORA requests; total pay only"),
    "ne": dict(state="NE", system="University of Nebraska budgeted employees", catalog_url_like="%nebraska.edu%", parse=parse_ne,
               notes="budgeted annual salary as of July 1; fiscal_year = ending year of the FY"),
    "me": dict(state="ME", system="University of Maine System regular employee compensation", catalog_url_like="%sharepoint%", parse=parse_me,
               notes="annual base salary snapshot June 2026 (FY2026)"),
    "nm": dict(state="NM", system="University of New Mexico Sunshine salary book", catalog_url_like="%unm.edu%", parse=parse_nm,
               notes="current snapshot (retrieved Oct 2026 = FY2027); primary position salary only"),
    "dc": dict(state="DC", system="DC Government employee salary layer (UDC rows)", catalog_url_like="%dcgis%", parse=parse_dc,
               notes="point-in-time compensation rate; hourly/per-course rates kept in pay_components only; fiscal_year = DC FY2027 (snapshot Oct 2026)"),
    "az_asu": dict(state="AZ", system="Arizona State University (State Press SalaryList)", catalog_url_like="%statepress%", parse=parse_az_asu,
                   notes="base salary by calendar year; revised files used where available"),
    "az_ua": dict(state="AZ", system="University of Arizona (Flourish embed)", catalog_url_like="%flo.uri.sh%", parse=parse_az_ua,
                  notes="salary at employment FTE by fiscal year; third-party newsroom"),
    "ok": dict(state="OK", system="Oklahoma OMES state payroll (higher-ed agencies), FY2026", catalog_url_like="%data.ok.gov%", parse=parse_ok,
               notes="monthly payment lines summed per agency + last name + initials per month (FY2026 files 2025-07..2026-06); names are initials only, so persons are ambiguous"),
    "nj": dict(state="NJ", system="NJ pension membership (active members, higher-ed employers)", catalog_url_like="%44xg-bswk%", parse=parse_nj,
               notes="pension-enrolled employees only (excludes ABP faculty); salary as reported to the pension system as of 2026-06-30; no title"),
    "il_graybook": dict(state="IL", system="University of Illinois Board of Trustees Gray Book 2026-27", catalog_url_like="%Graybook%", parse=parse_il_graybook,
                        file_filter=lambda f: f in ("UIUC.html", "UIC.html", "UIS.html", "SysOfc.html"),
                        notes="campus-level pages only (unit pages repeat the same rows); one row per job line as published (multi-job employees have several rows); base_pay = proposed FY2027 salary; present FY2026 salary in pay_components"),
    "sc": dict(state="SC", system="SC Dept of Administration state salaries (colleges and universities)", catalog_url_like="%admin.sc.gov%", parse=parse_sc,
               notes="current FY (last updated 2026-09-16 => FY2027); total compensation = annual salary or YTD earnings, whichever greater; employees under $50k omitted by the publisher"),
    "md": dict(state="MD", system="University of Maryland, College Park (Diamondback salary guide API)", catalog_url_like="%dbknews%", parse=parse_md,
               notes="student newspaper (public records); annual salary as listed for each year label; year label = fiscal year (verified against known "
                     "tenure dates, see parsers.parse_md); people on settlements stay listed while paid; 2025 names are 'First Last', earlier 'Last, First'"),
    "ne_cc": dict(state="NE", system="Nebraska community colleges (Flatwater Free Press Nebraska Public Payrolls)", catalog_url_like="%flatwaterfreepress%",
                  parse=parse_ne_cc,
                  notes="newsroom compile of public records; WordPress REST API (GET) for Metropolitan, Mid-Plains and Southeast CC; total pay as compiled by "
                        "the publisher; fiscal_year only where the publisher gives a year, else NULL (batch label in pay_components.ffp_update_batch)"),
    "mo_stl": dict(state="MO", system="St. Louis Post-Dispatch Public Pay 2026 (Missouri universities and colleges)", catalog_url_like="%stltoday%",
                   parse=parse_mo_stl_mo, data_key="mo_stl",
                   file_filter=lambda f: f in ("stl_115.html", "stl_116.html", "stl_117.html", "stl_118.html", "stl_119.html", "stl_120.html", "stl_78.html", "stl_87.html"),
                   notes="newsroom compile of public records; actual gross pay for calendar 2025 incl. overtime/additional pay; multiple jobs combined "
                         "under the highest-paid title; plain-GET HTML pages saved per agency"),
    "il_stl": dict(state="IL", system="St. Louis Post-Dispatch Public Pay 2026 (SIUE and University of Illinois)", catalog_url_like="%stltoday%",
                   parse=parse_mo_stl_il, data_key="mo_stl", catalog_state="MO",
                   file_filter=lambda f: f in ("stl_111.html", "stl_112.html", "stl_113.html", "stl_114.html", "stl_136.html"),
                   notes="same files as mo_stl; Illinois agencies (SIUE, UIUC, UIC, UIS, UI System) - actual gross pay for calendar 2025"),
}

ORDER = ["ct", "ma", "nyc", "ga", "tx", "il_iccb", "oh_osu", "oh_ou", "mn", "ri", "wa", "ne", "me", "nm", "dc", "az_asu", "az_ua", "ks", "nj", "ok", "il_graybook", "sc", "md", "ne_cc", "mo_stl", "il_stl"]


# ---------------------------------------------------------------- follow-up modules (PDF and other B sources)
import importlib as _il

for _mod in ("pdf_mo_il", "pdf_misc", "pdf_more"):
    try:
        _m = _il.import_module(_mod)
    except ModuleNotFoundError:
        continue
    for _k, _v in getattr(_m, "SPECS", {}).items():
        SPECS[_k] = _v
        if _k not in ORDER:
            ORDER.append(_k)
