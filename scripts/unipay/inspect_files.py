"""Print headers / first rows of downloaded files (dev helper)."""
import csv, io, sys, zipfile
from pathlib import Path
import openpyxl

for p in map(Path, sys.argv[1:]):
    print("=====", p.name, p.stat().st_size)
    if p.suffix == ".zip":
        z = zipfile.ZipFile(p)
        for n in z.namelist()[:30]:
            print("  member", n, z.getinfo(n).file_size)
        n = z.namelist()[0]
        with z.open(n) as f:
            for i, line in enumerate(io.TextIOWrapper(f, encoding="latin-1")):
                print("  ", line.rstrip()[:300])
                if i > 4: break
    elif p.suffix in (".csv", ".json", ".html", ".jsonl"):
        with open(p, encoding="utf-8", errors="replace") as f:
            for i, line in enumerate(f):
                print("  ", line.rstrip()[:400])
                if i > 4: break
    elif p.suffix == ".xlsx":
        wb = openpyxl.load_workbook(p, read_only=True)
        for ws in wb.worksheets:
            print("  sheet", ws.title, ws.max_row, ws.max_column)
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if any(v is not None for v in row):
                    print("    ", i, [v for v in row][:20])
                if i > 8: break
